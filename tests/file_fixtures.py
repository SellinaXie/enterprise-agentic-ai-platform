"""Small deterministic in-memory enterprise files for V6.5 tests."""

from io import BytesIO

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def build_pdf(*page_texts: str, encrypted: bool = False) -> bytes:
    """Create a minimal text PDF without network or external fixture files."""
    output = BytesIO()
    writer = PdfWriter()
    for text in page_texts:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        safe_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({safe_text}) Tj ET".encode("latin-1"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("synthetic-password")
    writer.write(output)
    return output.getvalue()


def build_docx(*, include_content: bool = True) -> bytes:
    """Create a DOCX with a heading, paragraphs, and a small table."""
    document = Document()
    if include_content:
        document.core_properties.title = "Synthetic Governance Manual"
        document.add_heading("Human Review", level=1)
        document.add_paragraph("Authorized reviewers approve high-impact recommendations.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Control"
        table.cell(0, 1).text = "Owner"
        table.cell(1, 0).text = "Manual approval"
        table.cell(1, 1).text = "Risk team"
    output = BytesIO()
    document.save(output)
    return output.getvalue()
